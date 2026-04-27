"""L1 单元测试 - fix-unit-price-orphan-fallback

覆盖 ``_find_pricing_xlsx`` 与 ``_find_all_pricing_xlsx`` 的 fallback 行为
+ 单 bidder 单类不变量 + leader/fill 对称性 + sheet mismatch regression。

详见 spec ``parser-pipeline`` ADDED Requirement
"报价 XLSX 选取 fallback 与单 bidder 单类不变量"。
"""

from __future__ import annotations

import pytest
import pytest_asyncio
from sqlalchemy import delete, select

from app.db.session import async_session
from app.models.bid_document import BidDocument
from app.models.bidder import Bidder
from app.models.price_item import PriceItem
from app.models.price_parsing_rule import PriceParsingRule
from app.models.project import Project
from app.models.user import User
from app.services.parser.pipeline.fill_price import fill_price_from_rule
from app.services.parser.pipeline.run_pipeline import (
    _find_all_pricing_xlsx,
    _find_pricing_xlsx,
)


_PREFIX = "pxf_"


async def _purge() -> None:
    async with async_session() as s:
        user_ids = (
            await s.execute(
                select(User.id).where(User.username.like(f"{_PREFIX}%"))
            )
        ).scalars().all()
        if not user_ids:
            return
        project_ids = (
            await s.execute(
                select(Project.id).where(Project.owner_id.in_(user_ids))
            )
        ).scalars().all()
        bidder_ids = (
            (
                await s.execute(
                    select(Bidder.id).where(Bidder.project_id.in_(project_ids))
                )
            ).scalars().all()
            if project_ids else []
        )
        if bidder_ids:
            await s.execute(
                delete(PriceItem).where(PriceItem.bidder_id.in_(bidder_ids))
            )
            await s.execute(
                delete(BidDocument).where(BidDocument.bidder_id.in_(bidder_ids))
            )
            await s.execute(
                delete(Bidder).where(Bidder.id.in_(bidder_ids))
            )
        if project_ids:
            await s.execute(
                delete(PriceParsingRule).where(
                    PriceParsingRule.project_id.in_(project_ids)
                )
            )
            await s.execute(delete(Project).where(Project.id.in_(project_ids)))
        await s.execute(delete(User).where(User.id.in_(user_ids)))
        await s.commit()


async def _seed_bidder(role_files: list[tuple[str, str]]) -> int:
    """seed 1 user/project/bidder + 给定 (file_role, file_name) 的 BidDocument 列表。

    所有 doc parse_status='identified' file_type='.xlsx'。
    返回 bidder_id。调用方负责测试结束后调 _purge。
    """
    async with async_session() as s:
        user = User(
            username=f"{_PREFIX}{id(role_files)}",
            password_hash="x",
            role="reviewer",
            login_fail_count=0,
        )
        s.add(user)
        await s.flush()
        project = Project(name=f"{_PREFIX}P", owner_id=user.id)
        s.add(project)
        await s.flush()
        bidder = Bidder(
            name=f"{_PREFIX}B",
            project_id=project.id,
            parse_status="identified",
        )
        s.add(bidder)
        await s.flush()
        for role, fname in role_files:
            doc = BidDocument(
                bidder_id=bidder.id,
                file_name=fname,
                file_path=f"/fake/{fname}",
                file_type=".xlsx",
                file_size=1024,
                md5=f"deadbeef{fname}",  # NOT NULL constraint
                source_archive=f"fake-{fname}.zip",  # NOT NULL constraint
                file_role=role,
                role_confidence="high",
                parse_status="identified",
            )
            s.add(doc)
        await s.commit()
        return bidder.id


@pytest_asyncio.fixture(autouse=True)
async def _cleanup():
    await _purge()
    yield
    await _purge()


# ---- 2.1 / 2.2 :4 个 Scenario × 2 个 helper ----


async def test_find_all_pricing_only_returns_pricing() -> None:
    """Scenario 1:仅 pricing 类 XLSX(主路径)。"""
    bidder_id = await _seed_bidder(
        [("pricing", "main.xlsx"), ("pricing", "main2.xlsx")]
    )
    paths = await _find_all_pricing_xlsx(bidder_id)
    assert sorted(paths) == ["/fake/main.xlsx", "/fake/main2.xlsx"]


async def test_find_all_unit_price_fallback() -> None:
    """Scenario 2:仅 unit_price 类 XLSX(fallback)。"""
    bidder_id = await _seed_bidder(
        [("unit_price", "u.xlsx"), ("unit_price", "u2.xlsx")]
    )
    paths = await _find_all_pricing_xlsx(bidder_id)
    assert sorted(paths) == ["/fake/u.xlsx", "/fake/u2.xlsx"]


async def test_find_all_both_present_prefers_pricing() -> None:
    """Scenario 3:pricing + unit_price 都有,**仅** pricing 进回填(不变量)。"""
    bidder_id = await _seed_bidder(
        [("pricing", "main.xlsx"), ("unit_price", "u.xlsx")]
    )
    paths = await _find_all_pricing_xlsx(bidder_id)
    # 关键不变量:返回的列表 **不** 包含 unit_price 类
    assert paths == ["/fake/main.xlsx"]
    assert "/fake/u.xlsx" not in paths


async def test_find_all_neither_returns_empty() -> None:
    """Scenario 4:既无 pricing 也无 unit_price → 空列表。"""
    bidder_id = await _seed_bidder(
        [("technical", "t.xlsx"), ("qualification", "q.xlsx")]
    )
    paths = await _find_all_pricing_xlsx(bidder_id)
    assert paths == []


async def test_find_pricing_only_returns_pricing() -> None:
    """leader 选举 Scenario 1:仅 pricing 类。"""
    bidder_id = await _seed_bidder(
        [("pricing", "main.xlsx"), ("pricing", "main2.xlsx")]
    )
    path = await _find_pricing_xlsx(bidder_id)
    assert path in ("/fake/main.xlsx", "/fake/main2.xlsx")


async def test_find_pricing_unit_price_fallback() -> None:
    """leader 选举 Scenario 2:仅 unit_price → fallback。"""
    bidder_id = await _seed_bidder([("unit_price", "u.xlsx")])
    path = await _find_pricing_xlsx(bidder_id)
    assert path == "/fake/u.xlsx"


async def test_find_pricing_both_present_prefers_pricing() -> None:
    """leader 选举 Scenario 3:两者都有 → 必然返 pricing。"""
    bidder_id = await _seed_bidder(
        [("pricing", "main.xlsx"), ("unit_price", "u.xlsx")]
    )
    path = await _find_pricing_xlsx(bidder_id)
    assert path == "/fake/main.xlsx"


async def test_find_pricing_neither_returns_none() -> None:
    """leader 选举 Scenario 4:都无 → None。"""
    bidder_id = await _seed_bidder(
        [("technical", "t.xlsx"), ("qualification", "q.xlsx")]
    )
    path = await _find_pricing_xlsx(bidder_id)
    assert path is None


# ---- 2.3:leader/fill 对称性测试(3 个 fixture)----


@pytest.mark.parametrize(
    "fixture_name,role_files,expected_role",
    [
        (
            "pure_pricing",
            [("pricing", "main.xlsx")],
            "pricing",
        ),
        (
            "pure_unit_price",
            [("unit_price", "u.xlsx")],
            "unit_price",
        ),
        (
            "both_prefer_pricing",
            [("pricing", "main.xlsx"), ("unit_price", "u.xlsx")],
            "pricing",
        ),
    ],
)
async def test_leader_and_fill_symmetric(
    fixture_name: str,
    role_files: list[tuple[str, str]],
    expected_role: str,
) -> None:
    """对称性回归保护:`_find_pricing_xlsx` 选出的 file_path
    必然属于 `_find_all_pricing_xlsx` 返回的列表;两者使用相同 role。

    防止未来有人改动其中一个 helper 忘改另一个,导致 leader 选 unit_price
    样本但回填走 pricing 路径(或反向)。
    """
    bidder_id = await _seed_bidder(role_files)

    leader_path = await _find_pricing_xlsx(bidder_id)
    all_paths = await _find_all_pricing_xlsx(bidder_id)

    assert leader_path is not None, f"{fixture_name}: leader 不应为 None"
    assert leader_path in all_paths, (
        f"{fixture_name}: leader 选的 {leader_path} 不在 all 列表 {all_paths}"
    )

    # 验证 all 返回的所有文件都是同一 role(单 bidder 单类不变量)
    async with async_session() as s:
        stmt = select(BidDocument).where(
            BidDocument.bidder_id == bidder_id,
            BidDocument.file_path.in_(all_paths),
        )
        docs = (await s.execute(stmt)).scalars().all()
    actual_roles = {d.file_role for d in docs}
    assert actual_roles == {expected_role}, (
        f"{fixture_name}: all 列表混合了 {actual_roles},应只含 {expected_role}"
    )


# ---- 2.4:sheet mismatch regression(fallback 命中但 sheet 严格匹配仍失败)----


async def test_unit_price_fallback_then_sheet_mismatch_fails(tmp_path) -> None:
    """fallback 选中 unit_price xlsx,但 LLM rule 的 sheet_name 严格不匹配
    → fill_price_from_rule 不抽出任何 price_items,partial_failed_sheets 含未找到。

    用意:fallback 只放宽"哪些 xlsx 入回填",**不**应放宽 sheet 名严格匹配。
    若未来有人误把 sheet 匹配也放宽,本 case 会拦住。
    """
    from openpyxl import Workbook

    # 构造 unit_price xlsx,内部 sheet 名 = "实际表名"
    xlsx_path = tmp_path / "u.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(title="实际表名")
    ws.append(["序号", "项目", "单价", "总价"])
    ws.append([1, "A", 100, 100])
    wb.save(xlsx_path)

    # seed bidder 仅含 1 份 unit_price xlsx
    async with async_session() as s:
        user = User(
            username=f"{_PREFIX}sheet_mismatch",
            password_hash="x",
            role="reviewer",
            login_fail_count=0,
        )
        s.add(user)
        await s.flush()
        project = Project(name=f"{_PREFIX}P_sm", owner_id=user.id)
        s.add(project)
        await s.flush()
        bidder = Bidder(
            name=f"{_PREFIX}B_sm",
            project_id=project.id,
            parse_status="pricing",
        )
        s.add(bidder)
        await s.flush()
        # rule 的 sheet_name 故意与 xlsx 内部不一致
        rule = PriceParsingRule(
            project_id=project.id,
            sheet_name="LLM 识别出的不存在表",
            header_row=1,
            column_mapping={
                "code_col": "A",
                "name_col": "B",
                "unit_price_col": "C",
                "total_price_col": "D",
            },
            sheets_config=[
                {
                    "sheet_name": "LLM 识别出的不存在表",
                    "header_row": 1,
                    "column_mapping": {
                        "code_col": "A",
                        "name_col": "B",
                        "unit_price_col": "C",
                        "total_price_col": "D",
                    },
                }
            ],
            status="confirmed",
            confirmed=True,
            created_by_llm=True,
        )
        s.add(rule)
        doc = BidDocument(
            bidder_id=bidder.id,
            file_name="u.xlsx",
            file_path=str(xlsx_path),
            file_type=".xlsx",
            file_size=xlsx_path.stat().st_size,
            md5="deadbeef_u_sm",
            source_archive="fake-u-sm.zip",
            file_role="unit_price",
            role_confidence="high",
            parse_status="identified",
        )
        s.add(doc)
        await s.commit()
        bidder_id = bidder.id
        rule_id = rule.id

    # 验证 fallback helper 选中 unit_price xlsx
    paths = await _find_all_pricing_xlsx(bidder_id)
    assert paths == [str(xlsx_path)], "fallback 应选中 unit_price xlsx"

    # 跑 fill_price_from_rule:sheet 名不匹配 → 0 行 + partial_failed
    async with async_session() as s:
        loaded_rule = await s.get(PriceParsingRule, rule_id)
        assert loaded_rule is not None
        result = await fill_price_from_rule(s, bidder_id, loaded_rule, paths[0])

    # 关键断言:fallback 命中文件,但 sheet 名不匹配仍然失败(不静默成功)
    assert result.items_count == 0, "sheet 名不匹配时不应回填任何行"
    assert any("未找到" in s for s in result.partial_failed_sheets), (
        f"应记录'未找到'失败信息,实际 partial_failed_sheets={result.partial_failed_sheets}"
    )
    # DB 验证无 price_items 残留
    async with async_session() as s:
        items = (
            await s.execute(
                select(PriceItem).where(PriceItem.bidder_id == bidder_id)
            )
        ).scalars().all()
    assert items == [], "sheet mismatch 不应有 price_items 入库"
